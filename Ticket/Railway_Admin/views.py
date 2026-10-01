import secrets
import string
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile

from django.http import HttpResponse, JsonResponse
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib.auth.forms import PasswordChangeForm
from django.core.mail import send_mail
from django.conf import settings
from django.db import transaction
from django.core.paginator import Paginator
from django.urls import reverse
from django.shortcuts import render, redirect
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.dateparse import parse_datetime
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from .models import TrainInformation, TrainSchedule, TrainDriverInformation
from django.contrib.auth.decorators import user_passes_test
from django.shortcuts import render, redirect, get_object_or_404
from .models import AdminProfile



def is_superuser(user):
    return user.is_authenticated and user.is_superuser


def superuser_dashboard(request):
    if not request.user.is_authenticated:
        return render(request, 'Railway_Admin/superuser_base.html', {
            'login_landing': True,
            'next': reverse('superuser_dashboard'),
        })
    if not request.user.is_superuser:
        if not hasattr(request.user, 'adminprofile'):
            return redirect('superuser_login')
        if not request.user.adminprofile.one_time_password:
            return redirect('admin_password_change')
        return render(
            request,
            'Railway_Admin/admin_dashboard.html',
            _admin_dashboard_context(request),
        )

    profiles = AdminProfile.objects.select_related('user').all()
    return render(request, 'Railway_Admin/superuser_dashboard.html', {'profiles': profiles})


def superuser_login(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = User.objects.filter(username=username).first()

        if user and user.check_password(password) and (user.is_superuser or hasattr(user, 'adminprofile')):
            login(request, user)
            requested_next = request.POST.get('next', '')
            if user.is_superuser:
                redirect_url = (
                    requested_next
                    if url_has_allowed_host_and_scheme(
                        requested_next,
                        allowed_hosts={request.get_host()},
                        require_https=request.is_secure(),
                    )
                    else reverse('superuser_dashboard')
                )
            elif user.adminprofile.one_time_password:
                redirect_url = (
                    requested_next
                    if requested_next == reverse('superuser_dashboard')
                    else reverse('admin_dashboard')
                )
            else:
                redirect_url = reverse('admin_password_change')
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'redirect_url': redirect_url})
            return redirect(redirect_url)

        error = 'Invalid credentials or user account.'
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': error})
        return render(request, 'Railway_Admin/superuser_base.html', {
            'error': error,
            'next': request.POST.get('next', ''),
            'open_login': True,
            'login_landing': request.POST.get('next') == reverse('superuser_dashboard'),
        })

    return render(request, 'Railway_Admin/superuser_base.html', {
        'next': request.GET.get('next', ''),
        'open_login': True,
    })


def superuser_logout(request):
    logout(request)
    return redirect('superuser_login')


@login_required(login_url='superuser_login')
def admin_password_change(request):
    if request.user.is_superuser or not hasattr(request.user, 'adminprofile'):
        return redirect('admin_dashboard')

    profile = request.user.adminprofile
    if profile.one_time_password:
        return redirect('admin_dashboard')

    if request.method == 'POST':
        form = PasswordChangeForm(request.user, request.POST)
        if form.is_valid():
            form.save()
            profile.one_time_password = True
            profile.save(update_fields=['one_time_password'])
            login(request, request.user)
            return redirect('admin_dashboard')
    else:
        form = PasswordChangeForm(request.user)

    return render(request, 'Railway_Admin/admin_password_change.html', {'form': form})


@user_passes_test(is_superuser, login_url='superuser_login')
def admin_profile_create(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        full_name = request.POST.get('full_name')
        email = request.POST.get('email')
        designation = request.POST.get('designation')
        phone_number = request.POST.get('phone_number')

        if User.objects.filter(username=username).exists():
            return JsonResponse({'success': False, 'error': 'Username already exists.'})
        if AdminProfile.objects.filter(email=email).exists():
            return JsonResponse({'success': False, 'error': 'Email already exists.'})
        if AdminProfile.objects.filter(phone_number=phone_number).exists():
            return JsonResponse({'success': False, 'error': 'Phone number already exists.'})

        temporary_password = ''.join(
            secrets.choice(string.ascii_letters + string.digits)
            for _ in range(12)
        )

        try:
            with transaction.atomic():
                user = User.objects.create_user(
                    username=username,
                    email=email,
                    password=temporary_password,
                )
                AdminProfile.objects.create(
                    user=user, full_name=full_name, email=email,
                    designation=designation, phone_number=phone_number
                )
                send_mail(
                    'Your Railway Admin Account',
                    (
                        f'Hello {full_name},\n\n'
                        f'Your admin account has been created.\n'
                        f'Username: {username}\n'
                        f'One-time password: {temporary_password}\n\n'
                        'Please log in and change this password immediately.'
                    ),
                    settings.DEFAULT_FROM_EMAIL,
                    [email],
                    fail_silently=False,
                )
        except Exception:
            return JsonResponse({
                'success': False,
                'error': 'Admin profile could not be created or the email could not be sent.',
            }, status=500)

        return JsonResponse({'success': True})

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_superuser, login_url='superuser_login')
def admin_profile_get(request, pk):
    profile = get_object_or_404(AdminProfile, pk=pk)
    return JsonResponse({
        'id': profile.pk,
        'username': profile.user.username,
        'full_name': profile.full_name,
        'email': profile.email,
        'designation': profile.designation,
        'phone_number': profile.phone_number,
    })


@user_passes_test(is_superuser, login_url='superuser_login')
def admin_profile_edit(request, pk):
    profile = get_object_or_404(AdminProfile, pk=pk)

    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        if User.objects.filter(username=username).exclude(pk=profile.user.pk).exists():
            return JsonResponse({'success': False, 'error': 'Username already exists.'})

        profile.user.username = username
        if password:
            profile.user.set_password(password)
        profile.user.save()

        profile.full_name = request.POST.get('full_name')
        profile.email = request.POST.get('email')
        profile.designation = request.POST.get('designation')
        profile.phone_number = request.POST.get('phone_number')
        profile.save()

        return JsonResponse({'success': True})

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_superuser, login_url='superuser_login')
def admin_profile_delete(request, pk):
    profile = get_object_or_404(AdminProfile, pk=pk)
    if request.method == 'POST':
        profile.user.delete()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


def is_admin(user):
    return user.is_authenticated and hasattr(user, 'adminprofile')


def _admin_dashboard_context(request):
    trains = TrainInformation.objects.order_by('train_number')
    schedules = TrainSchedule.objects.select_related('train').order_by('departure_time', 'pk')
    return {
        'trains': trains,
        'train_page': Paginator(trains, 10).get_page(request.GET.get('train_page')),
        'schedules': schedules,
        'schedule_page': Paginator(schedules, 10).get_page(request.GET.get('schedule_page')),
        'drivers': TrainDriverInformation.objects.select_related('train').all(),
    }


@user_passes_test(is_admin, login_url='superuser_login')
def admin_dashboard(request):
    if not request.user.adminprofile.one_time_password:
        return redirect('admin_password_change')
    return render(request, 'Railway_Admin/admin_dashboard.html', _admin_dashboard_context(request))


@user_passes_test(is_admin, login_url='superuser_login')
@require_GET
def train_export(request):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Trains'
    worksheet.append(['Train Number', 'Train Name', 'Total Seats'])

    for train in TrainInformation.objects.order_by('train_number'):
        worksheet.append([train.train_number, train.train_name, train.total_seats])

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="trains.xlsx"'
    return response


@user_passes_test(is_admin, login_url='superuser_login')
@require_POST
def train_import(request):
    uploaded_file = request.FILES.get('file')
    if not uploaded_file or not uploaded_file.name.lower().endswith('.xlsx'):
        return JsonResponse({'success': False, 'error': 'Choose an .xlsx Excel file.'}, status=400)

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return JsonResponse({'success': False, 'error': 'The Excel sheet is empty.'}, status=400)

        header_indexes = {
            str(value).strip().lower(): index
            for index, value in enumerate(headers)
            if value is not None
        }
        required_headers = ('train number', 'train name', 'total seats')
        if any(header not in header_indexes for header in required_headers):
            return JsonResponse({
                'success': False,
                'error': 'Required headers: Train Number, Train Name, Total Seats.',
            }, status=400)

        valid_trains = []
        errors = []
        seen_train_numbers = set()
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue

            def value_for(header):
                index = header_indexes[header]
                return row[index] if index < len(row) else None

            raw_number = value_for('train number')
            if isinstance(raw_number, float) and raw_number.is_integer():
                train_number = str(int(raw_number))
            else:
                train_number = str(raw_number or '').strip()
            train_name = str(value_for('train name') or '').strip()
            try:
                seats_value = float(value_for('total seats'))
                total_seats = int(seats_value)
                if not seats_value.is_integer() or total_seats <= 0:
                    raise ValueError
            except (TypeError, ValueError, OverflowError):
                total_seats = 0

            if not train_number or not train_name or total_seats <= 0:
                errors.append(f'Row {row_number}: train number, train name, and positive whole-number seats are required.')
            elif train_number in seen_train_numbers:
                errors.append(f'Row {row_number}: duplicate train number {train_number}.')
            else:
                seen_train_numbers.add(train_number)
                valid_trains.append({
                    'train_number': train_number,
                    'train_name': train_name,
                    'total_seats': total_seats,
                })
        workbook.close()
    except (InvalidFileException, BadZipFile, OSError, ValueError, IndexError):
        return JsonResponse({'success': False, 'error': 'Could not read this Excel file.'}, status=400)

    if errors:
        return JsonResponse({'success': False, 'error': ' '.join(errors[:10])}, status=400)
    if not valid_trains:
        return JsonResponse({'success': False, 'error': 'The Excel sheet contains no train rows.'}, status=400)

    created_count = 0
    updated_count = 0
    with transaction.atomic():
        for train_data in valid_trains:
            _, created = TrainInformation.objects.update_or_create(
                train_number=train_data['train_number'],
                defaults={
                    'train_name': train_data['train_name'],
                    'total_seats': train_data['total_seats'],
                },
            )
            if created:
                created_count += 1
            else:
                updated_count += 1

    return JsonResponse({
        'success': True,
        'created': created_count,
        'updated': updated_count,
    })


@user_passes_test(is_admin, login_url='superuser_login')
@require_GET
def schedule_export(request):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Schedules'
    worksheet.append([
        'Schedule ID', 'Train Number', 'Source Station', 'Destination Station',
        'Departure Time', 'Arrival Time', 'AC Ticket Price',
        'Snigdha Ticket Price', 'S. Chair Ticket Price',
    ])
    worksheet.freeze_panes = 'A2'
    for schedule in TrainSchedule.objects.select_related('train').order_by('departure_time', 'pk'):
        worksheet.append([
            schedule.pk,
            schedule.train.train_number,
            schedule.source_station,
            schedule.destination_station,
            timezone.localtime(schedule.departure_time).replace(tzinfo=None),
            timezone.localtime(schedule.arrival_time).replace(tzinfo=None),
            schedule.ac_ticket_price,
            schedule.singdha_ticket_price,
            schedule.s_chair_ticket_price,
        ])

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="train-schedules.xlsx"'
    return response


@user_passes_test(is_admin, login_url='superuser_login')
@require_POST
def schedule_import(request):
    uploaded_file = request.FILES.get('file')
    if not uploaded_file or not uploaded_file.name.lower().endswith('.xlsx'):
        return JsonResponse({'success': False, 'error': 'Choose an .xlsx Excel file.'}, status=400)

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            workbook.close()
            return JsonResponse({'success': False, 'error': 'The Excel sheet is empty.'}, status=400)

        header_indexes = {
            str(value).strip().lower(): index
            for index, value in enumerate(headers)
            if value is not None
        }
        required_headers = (
            'train number', 'source station', 'destination station',
            'departure time', 'arrival time',
        )
        if any(header not in header_indexes for header in required_headers):
            workbook.close()
            return JsonResponse({
                'success': False,
                'error': 'Required headers: Train Number, Source Station, Destination Station, Departure Time, Arrival Time.',
            }, status=400)

        trains_by_number = {
            train.train_number: train
            for train in TrainInformation.objects.all()
        }
        valid_schedules = []
        errors = []
        seen_schedule_ids = set()

        def cell_value(row, header):
            index = header_indexes.get(header)
            return row[index] if index is not None and index < len(row) else None

        def parse_schedule_datetime(value):
            if isinstance(value, datetime):
                parsed = value
            elif isinstance(value, date):
                parsed = datetime.combine(value, datetime.min.time())
            else:
                parsed = parse_datetime(str(value or '').strip())
            if parsed is None:
                raise ValueError
            if settings.USE_TZ and timezone.is_naive(parsed):
                parsed = timezone.make_aware(parsed)
            return parsed

        def parse_schedule_price(value):
            if value is None or str(value).strip() == '':
                return None
            numeric = float(value)
            if not numeric.is_integer() or numeric < 0:
                raise ValueError
            return int(numeric)

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue

            raw_train_number = cell_value(row, 'train number')
            if isinstance(raw_train_number, float) and raw_train_number.is_integer():
                train_number = str(int(raw_train_number))
            else:
                train_number = str(raw_train_number or '').strip()
            train = trains_by_number.get(train_number)
            source_station = str(cell_value(row, 'source station') or '').strip()
            destination_station = str(cell_value(row, 'destination station') or '').strip()
            try:
                departure_time = parse_schedule_datetime(cell_value(row, 'departure time'))
                arrival_time = parse_schedule_datetime(cell_value(row, 'arrival time'))
                if arrival_time <= departure_time:
                    raise ValueError
                prices = {
                    'ac_ticket_price': parse_schedule_price(cell_value(row, 'ac ticket price')),
                    'singdha_ticket_price': parse_schedule_price(cell_value(row, 'snigdha ticket price')),
                    's_chair_ticket_price': parse_schedule_price(cell_value(row, 's. chair ticket price')),
                }
            except (TypeError, ValueError, OverflowError):
                errors.append(f'Row {row_number}: enter valid departure/arrival times and non-negative whole-number fares.')
                continue

            raw_schedule_id = cell_value(row, 'schedule id')
            try:
                schedule_id = int(raw_schedule_id) if raw_schedule_id not in (None, '') else None
                if schedule_id is not None and (float(raw_schedule_id) != schedule_id or schedule_id <= 0):
                    raise ValueError
            except (TypeError, ValueError, OverflowError):
                errors.append(f'Row {row_number}: Schedule ID must be a positive whole number or blank.')
                continue

            if not train:
                errors.append(f'Row {row_number}: train number {train_number or "(blank)"} was not found.')
            elif not source_station or not destination_station:
                errors.append(f'Row {row_number}: source and destination stations are required.')
            elif schedule_id is not None and schedule_id in seen_schedule_ids:
                errors.append(f'Row {row_number}: duplicate Schedule ID {schedule_id}.')
            else:
                if schedule_id is not None:
                    seen_schedule_ids.add(schedule_id)
                valid_schedules.append({
                    'id': schedule_id,
                    'train': train,
                    'source_station': source_station,
                    'destination_station': destination_station,
                    'departure_time': departure_time,
                    'arrival_time': arrival_time,
                    **prices,
                })
        workbook.close()
    except (InvalidFileException, BadZipFile, OSError, ValueError, IndexError):
        return JsonResponse({'success': False, 'error': 'Could not read this Excel file.'}, status=400)

    existing_schedule_ids = set(TrainSchedule.objects.filter(
        pk__in=[row['id'] for row in valid_schedules if row['id'] is not None]
    ).values_list('pk', flat=True))
    for row in valid_schedules:
        if row['id'] is not None and row['id'] not in existing_schedule_ids:
            errors.append(f'Schedule ID {row["id"]} was not found.')

    if errors:
        return JsonResponse({'success': False, 'error': ' '.join(errors[:10])}, status=400)
    if not valid_schedules:
        return JsonResponse({'success': False, 'error': 'The Excel sheet contains no schedule rows.'}, status=400)

    created_count = 0
    updated_count = 0
    with transaction.atomic():
        for row in valid_schedules:
            schedule_id = row.pop('id')
            if schedule_id is None:
                TrainSchedule.objects.create(**row)
                created_count += 1
            else:
                TrainSchedule.objects.filter(pk=schedule_id).update(**row)
                updated_count += 1

    return JsonResponse({
        'success': True,
        'created': created_count,
        'updated': updated_count,
    })


@user_passes_test(is_admin, login_url='superuser_login')
def train_create(request):
    if request.method == 'POST':
        train_number = request.POST.get('train_number')
        train_name = request.POST.get('train_name')
        total_seats = request.POST.get('total_seats')

        if TrainInformation.objects.filter(train_number=train_number).exists():
            return JsonResponse({'success': False, 'error': 'Train number already exists.'})

        train = TrainInformation.objects.create(
            train_number=train_number, train_name=train_name, total_seats=total_seats
        )
        return JsonResponse({
            'success': True,
            'id': train.pk,
            'train_number': train.train_number,
            'train_name': train.train_name,
            'total_seats': train.total_seats,
        })

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def train_get(request, pk):
    train = get_object_or_404(TrainInformation, pk=pk)
    return JsonResponse({
        'id': train.pk,
        'train_number': train.train_number,
        'train_name': train.train_name,
        'total_seats': train.total_seats,
    })


@user_passes_test(is_admin, login_url='superuser_login')
def train_edit(request, pk):
    train = get_object_or_404(TrainInformation, pk=pk)
    if request.method == 'POST':
        train_number = request.POST.get('train_number')
        if TrainInformation.objects.filter(train_number=train_number).exclude(pk=pk).exists():
            return JsonResponse({'success': False, 'error': 'Train number already exists.'})
        train.train_number = train_number
        train.train_name = request.POST.get('train_name')
        train.total_seats = request.POST.get('total_seats')
        train.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def train_delete(request, pk):
    if request.method == 'POST':
        get_object_or_404(TrainInformation, pk=pk).delete()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def schedule_create(request):
    if request.method == 'POST':
        departure_time = parse_datetime(request.POST.get('departure_time', ''))
        arrival_time = parse_datetime(request.POST.get('arrival_time', ''))

        if not departure_time or not arrival_time:
            return JsonResponse({
                'success': False,
                'error': 'Valid departure and arrival times are required.',
            })

        if timezone.is_naive(departure_time):
            departure_time = timezone.make_aware(departure_time)
        if timezone.is_naive(arrival_time):
            arrival_time = timezone.make_aware(arrival_time)

        def parse_price(field):
            val = request.POST.get(field, '').strip()
            return int(val) if val else None

        TrainSchedule.objects.create(
            train_id=request.POST.get('train'),
            departure_time=departure_time,
            arrival_time=arrival_time,
            source_station=request.POST.get('source_station'),
            destination_station=request.POST.get('destination_station'),
            ac_ticket_price=parse_price('ac_ticket_price'),
            singdha_ticket_price=parse_price('singdha_ticket_price'),
            s_chair_ticket_price=parse_price('s_chair_ticket_price'),
        )
        return JsonResponse({'success': True})

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def schedule_get(request, pk):
    schedule = get_object_or_404(TrainSchedule, pk=pk)
    return JsonResponse({
        'id': schedule.pk,
        'train': schedule.train_id,
        'source_station': schedule.source_station,
        'destination_station': schedule.destination_station,
        'departure_time': timezone.localtime(schedule.departure_time).strftime('%Y-%m-%dT%H:%M'),
        'arrival_time': timezone.localtime(schedule.arrival_time).strftime('%Y-%m-%dT%H:%M'),
    })


@user_passes_test(is_admin, login_url='superuser_login')
def schedule_edit(request, pk):
    schedule = get_object_or_404(TrainSchedule, pk=pk)
    if request.method == 'POST':
        departure_time = parse_datetime(request.POST.get('departure_time', ''))
        arrival_time = parse_datetime(request.POST.get('arrival_time', ''))
        if not departure_time or not arrival_time:
            return JsonResponse({'success': False, 'error': 'Valid departure and arrival times are required.'})
        schedule.train_id = request.POST.get('train')
        schedule.source_station = request.POST.get('source_station')
        schedule.destination_station = request.POST.get('destination_station')
        schedule.departure_time = timezone.make_aware(departure_time) if timezone.is_naive(departure_time) else departure_time
        schedule.arrival_time = timezone.make_aware(arrival_time) if timezone.is_naive(arrival_time) else arrival_time
        schedule.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def schedule_delete(request, pk):
    if request.method == 'POST':
        get_object_or_404(TrainSchedule, pk=pk).delete()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def driver_create(request):
    if request.method == 'POST':
        license_number = request.POST.get('license_number')

        if TrainDriverInformation.objects.filter(license_number=license_number).exists():
            return JsonResponse({'success': False, 'error': 'License number already exists.'})

        TrainDriverInformation.objects.create(
            train_id=request.POST.get('train'),
            driver_name=request.POST.get('driver_name'),
            license_number=license_number,
            contact_number=request.POST.get('contact_number'),
        )
        return JsonResponse({'success': True})

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def driver_get(request, pk):
    driver = get_object_or_404(TrainDriverInformation, pk=pk)
    return JsonResponse({
        'id': driver.pk,
        'train': driver.train_id,
        'driver_name': driver.driver_name,
        'license_number': driver.license_number,
        'contact_number': driver.contact_number,
    })


@user_passes_test(is_admin, login_url='superuser_login')
def driver_edit(request, pk):
    driver = get_object_or_404(TrainDriverInformation, pk=pk)
    if request.method == 'POST':
        license_number = request.POST.get('license_number')
        if TrainDriverInformation.objects.filter(license_number=license_number).exclude(pk=pk).exists():
            return JsonResponse({'success': False, 'error': 'License number already exists.'})
        driver.train_id = request.POST.get('train')
        driver.driver_name = request.POST.get('driver_name')
        driver.license_number = license_number
        driver.contact_number = request.POST.get('contact_number')
        driver.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})


@user_passes_test(is_admin, login_url='superuser_login')
def driver_delete(request, pk):
    if request.method == 'POST':
        get_object_or_404(TrainDriverInformation, pk=pk).delete()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False, 'error': 'Invalid request.'})

@user_passes_test(is_admin, login_url='superuser_login')
def admin_profile(request):
    profile = request.user.adminprofile
    return render(request, 'Railway_Admin/admin_profile.html', {'profile': profile})


@user_passes_test(is_admin, login_url='superuser_login')
def admin_profile_update(request):
    profile = request.user.adminprofile

    if request.method == 'POST':
        full_name = request.POST.get('full_name')
        email = request.POST.get('email')
        designation = request.POST.get('designation')
        phone_number = request.POST.get('phone_number')

        current_password = request.POST.get('current_password')
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')

        # Email/phone uniqueness check (exclude nijer profile)
        if AdminProfile.objects.filter(email=email).exclude(pk=profile.pk).exists():
            return JsonResponse({'success': False, 'error': 'Email already in use.'})
        if AdminProfile.objects.filter(phone_number=phone_number).exclude(pk=profile.pk).exists():
            return JsonResponse({'success': False, 'error': 'Phone number already in use.'})

        # Password change (optional)
        if new_password or current_password or confirm_password:
            if not request.user.check_password(current_password):
                return JsonResponse({'success': False, 'error': 'Current password is incorrect.'})
            if new_password != confirm_password:
                return JsonResponse({'success': False, 'error': 'New passwords do not match.'})
            if len(new_password) < 8:
                return JsonResponse({'success': False, 'error': 'Password must be at least 8 characters.'})

            request.user.set_password(new_password)
            request.user.save()

        profile.full_name = full_name
        profile.email = email
        profile.designation = designation
        profile.phone_number = phone_number
        profile.save()

        # Password change hoile session invalidate hoye jay, tai re-login lagbe
        if new_password:
            return JsonResponse({'success': True, 'password_changed': True, 'redirect_url': reverse('superuser_login')})

        return JsonResponse({'success': True, 'password_changed': False})

    return JsonResponse({'success': False, 'error': 'Invalid request.'})


def superuser_logout(request):
    logout(request)
    return redirect('superuser_dashboard')


